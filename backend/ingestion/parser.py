import csv, io, hashlib, zipfile
from pathlib import Path
from openpyxl import load_workbook
import pdfplumber
from backend.storage.models import (
    Document,
    Span,
    Requirement,
    Proposal,
    Line,
    Catalog,
    Offer,
    uid,
)
from backend.storage.db import BLOBS
from backend.domain.service import get, scoped, lines, validate_sources, fail
from backend.domain.schemas import LineInput, Operation

ALLOWED = {".pdf", ".csv", ".xlsx", ".txt", ".eml"}
FIELDS = (
    "tag",
    "description",
    "quantity",
    "unit",
    "model",
    "cost",
    "price",
    "lead_time",
    "manufacturer",
    "supplier",
    "valid_until",
    "currency",
    "price_basis",
)
ALIASES = {
    "tag": ["tag", "equipment tag", "equipment_tag", "mark", "item"],
    "description": ["description", "equipment", "name"],
    "quantity": ["quantity", "qty", "count"],
    "unit": ["unit", "uom"],
    "model": ["model", "model number", "catalog"],
    "cost": ["cost", "unit cost", "supplier cost"],
    "price": ["price", "selling price", "unit price"],
    "lead_time": ["lead_time", "lead time", "delivery"],
    "manufacturer": ["manufacturer", "brand"],
    "supplier": ["supplier", "vendor"],
    "valid_until": ["valid_until", "valid until", "expiry"],
    "currency": ["currency", "ccy"],
    "price_basis": ["price_basis", "price basis", "price unit"],
}


def persist_blob(data, suffix):
    key = uid() + suffix
    (BLOBS / key).write_bytes(data)
    return key, hashlib.sha256(data).hexdigest()


def parse(s, d):
    path = BLOBS / d.blob
    ext = path.suffix.lower()
    spans = []
    coverage = []
    headers = []
    rows = []

    def span(text, location):
        sp = Span(id=uid(), document_id=d.id, text=text, location=location)
        s.add(sp)
        spans.append(sp)
        return sp

    if ext == ".eml":
        from email import policy
        from email.parser import BytesParser

        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        part = message.get_body(preferencelist=("plain",))
        body = part.get_content() if part else ""
        d.meta = {
            **d.meta,
            "email": {
                "subject": str(message.get("Subject", "")),
                "from": str(message.get("From", "")),
                "to": str(message.get("To", "")),
                "date": str(message.get("Date", "")),
                "message_id": str(message.get("Message-ID", "")),
            },
        }
        if body.strip():
            span(
                body.strip(),
                {
                    "label": "Email body",
                    "author": str(message.get("From", "")),
                    "authored_at": str(message.get("Date", "")),
                    "email": True,
                },
            )
        coverage = [
            {
                "label": "Email",
                "state": "processed" if body.strip() else "unsupported",
                "detail": "Plain-text email body; attachments are imported separately.",
            }
        ]
    elif ext == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        for n, row in enumerate(reader, 2):
            if n > 10002:
                raise ValueError("Schedules are limited to 10,000 rows per file.")
            values = {str(k): str(v or "") for k, v in row.items() if k is not None}
            sp = span(
                " | ".join(f"{k}: {v}" for k, v in values.items()),
                {
                    "sheet": "CSV",
                    "row": n,
                    "cells": f"A{n}:{chr(64 + min(len(headers), 26))}{n}",
                },
            )
            rows.append({"values": values, "source_id": sp.id})
        coverage = [{"label": "CSV", "state": "processed", "rows": len(rows)}]
    elif ext == ".xlsx":
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 80 MB.")
        formulas = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        for sheet in formulas:
            sheet_headers = []
            count = 0
            missing = False
            cached_rows = cached[sheet.title].iter_rows()
            for n, (row, cache) in enumerate(zip(sheet.iter_rows(), cached_rows), 1):
                if n > 10002:
                    raise ValueError("Sheets are limited to 10,000 rows.")
                vals = []
                formula_values = {}
                for cell, cache_cell in zip(row, cache):
                    if cell.data_type == "f":
                        formula_values[cell.coordinate] = str(cell.value)
                        if cache_cell.value is None:
                            missing = True
                        vals.append(
                            "" if cache_cell.value is None else str(cache_cell.value)
                        )
                    else:
                        vals.append("" if cell.value is None else str(cell.value))
                if not any(vals) and not formula_values:
                    continue
                if not sheet_headers:
                    sheet_headers = vals
                    headers = list(dict.fromkeys(headers + vals))
                    continue
                values = dict(zip(sheet_headers, vals))
                sp = span(
                    " | ".join(f"{k}: {v}" for k, v in values.items()),
                    {
                        "sheet": sheet.title,
                        "row": n,
                        "cells": f"A{n}:{row[-1].column_letter}{n}",
                        "formulas": formula_values,
                    },
                )
                rows.append({"values": values, "source_id": sp.id})
                count += 1
            coverage.append(
                {
                    "label": sheet.title,
                    "state": "unsupported" if missing else "processed",
                    "rows": count,
                    "detail": "Formula cells without cached values need review."
                    if missing
                    else "",
                }
            )
        formulas.close()
        cached.close()
    elif ext == ".pdf":
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 150:
                raise ValueError("PDFs are limited to 150 pages.")
            for n, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    unreadable = bool(page.images or page.curves)
                    page_source = (
                        span(
                            "",
                            {
                                "page": n,
                                "bbox": [0, 0, 1, 1],
                                "approximate": True,
                                "unreadable": True,
                            },
                        )
                        if unreadable
                        else None
                    )
                    coverage.append(
                        {
                            "label": f"Page {n}",
                            "state": "scanned" if unreadable else "empty",
                            "source_id": page_source.id if page_source else None,
                            "detail": "No extractable text. Transcription or visual review required.",
                        }
                    )
                    continue
                words = page.extract_words()
                groups = {}
                for word in words:
                    groups.setdefault(round(word["top"] / 8), []).append(word)
                for group in groups.values():
                    span(
                        " ".join(w["text"] for w in group),
                        {
                            "page": n,
                            "bbox": [
                                min(w["x0"] for w in group) / page.width,
                                min(w["top"] for w in group) / page.height,
                                max(w["x1"] for w in group) / page.width,
                                max(w["bottom"] for w in group) / page.height,
                            ],
                            "rotation": page.rotation,
                            "approximate": True,
                        },
                    )
                coverage.append(
                    {"label": f"Page {n}", "state": "processed", "words": len(words)}
                )
        # PDF annotation objects carry review comments separately from page text.
        from backend.orders.extraction import parse_pdf_annotations

        for annotation in parse_pdf_annotations(path):
            span(
                annotation["text"],
                {
                    "page": annotation["page"],
                    "bbox": annotation["bbox"],
                    "annotation_id": annotation["id"],
                    "annotation_type": annotation["kind"],
                    "author": annotation.get("author", ""),
                    "authored_at": annotation.get("authored_at", ""),
                    "approximate": False,
                },
            )
            if annotation.get("needs_review"):
                coverage.append(
                    {
                        "label": f"Page {annotation['page']} annotation",
                        "annotation_id": annotation["id"],
                        "state": "unsupported",
                        "detail": "A graphical annotation has no extractable text. Manual interpretation is required.",
                    }
                )
    else:
        text = path.read_text(encoding="utf-8-sig")
        for n, line in enumerate(text.splitlines(), 1):
            if line.strip():
                span(line, {"line": n, "label": f"Line {n}"})
        coverage = [
            {"label": "Text", "state": "processed" if text.strip() else "empty"}
        ]
    suggested = {
        field: next((h for h in headers if h.strip().lower() in names), "")
        for field, names in ALIASES.items()
    }
    d.coverage = coverage
    d.meta = {
        **d.meta,
        "headers": headers,
        "rows": rows,
        "suggested_mapping": suggested,
        "parser_version": "rivet-1",
    }
    d.state = "needs_mapping" if rows else "ready"
    s.flush()


def map_document(s, d, q, p, mapping, purpose):
    rows = d.meta.get("rows", [])
    if not rows:
        fail(
            "This source has no spreadsheet rows. Use the assistant or enter lines manually."
        )
    unknown = set(mapping) - set(FIELDS)
    if unknown:
        fail("Unknown mapping fields: " + ", ".join(unknown))
    if any(v and v not in d.meta.get("headers", []) for v in mapping.values()):
        fail("The selected column does not exist.")
    required = (
        ["model", "description", "manufacturer"]
        if purpose == "catalog"
        else ["model", "cost", "supplier", "unit", "valid_until"]
        if purpose == "offers"
        else ["tag", "quantity"]
    )
    if any(not mapping.get(k) for k in required):
        fail("Map the required columns: " + ", ".join(required))
    operations = []
    seen = {}
    issues = []
    bytag = {l.tag: l for l in lines(s, q)}
    for row in rows:
        data = {
            f: str(row["values"].get(h, "")).strip() for f, h in mapping.items() if h
        }
        # Never silently label a non-USD amount as USD or convert a price basis.
        for header, value in row["values"].items():
            if header.strip().lower() in ALIASES["currency"] and str(
                value
            ).strip().upper() not in ("", "USD"):
                fail(
                    "Only USD is supported. Resolve the source currency before importing."
                )
            if (
                header.strip().lower() in ALIASES["price_basis"]
                and str(value).strip()
                and str(value).strip().lower() != (data.get("unit") or "each").lower()
            ):
                fail(
                    "The supplier price basis differs from the quantity unit. A reviewed conversion is required."
                )
        source = row["source_id"]
        if purpose == "catalog":
            if not data.get("model"):
                issues.append("A catalog row has no model.")
                continue
            s.add(
                Catalog(
                    manufacturer=data.get("manufacturer", ""),
                    model=data["model"],
                    description=data.get("description", ""),
                    attributes={"source_ids": [source]},
                )
            )
            continue
        if purpose == "offers":
            item = s.scalar(scoped(Catalog).where(Catalog.model == data.get("model")))
            if not item:
                issues.append(f"Model {data.get('model')} is absent from the catalog.")
                continue
            from backend.domain.pricing import dec, money
            from datetime import date

            cost = dec(data.get("cost"))
            if cost < 0 or cost != money(cost):
                fail("Offer cost must be a nonnegative amount in cents.")
            date.fromisoformat(data["valid_until"])
            if data.get("unit") not in ("each", "ft", "package"):
                fail("Offer unit must be each, ft, or package.")
            s.add(
                Offer(
                    catalog_id=item.id,
                    supplier=data.get("supplier", ""),
                    cost=cost,
                    unit=data["unit"],
                    valid_until=data["valid_until"],
                    lead_time=data.get("lead_time", "Needs confirmation"),
                    sources=[source],
                )
            )
            continue
        tag = data.get("tag", "")
        if not tag:
            issues.append("A row is missing an equipment tag.")
            continue
        if tag in seen:
            if seen[tag] != data:
                issues.append(f"{tag}: conflicting duplicate rows require review.")
            continue
        seen[tag] = data
        if not data.get("quantity"):
            issues.append(f"{tag}: quantity is missing.")
            continue
        for prior in s.scalars(
            scoped(Requirement).where(
                Requirement.project_id == p.id,
                Requirement.tag == tag,
                Requirement.attribute == "quantity",
                Requirement.active == True,
            )
        ):
            if (
                purpose != "addendum"
                and prior.value.get("document_id") != d.id
                and prior.value.get("quantity") != data["quantity"]
            ):
                issues.append(
                    f"{tag}: conflicting requirements. Resolve precedence by identifying an authorized addendum."
                )
                continue
            prior.active = False
        s.add(
            Requirement(
                project_id=p.id,
                tag=tag,
                attribute="quantity",
                value={
                    "quantity": data["quantity"],
                    "unit": data.get("unit") or "each",
                    "document_id": d.id,
                },
                sources=[source],
                input_revision=p.input_revision,
            )
        )
        if tag in bytag:
            l = bytag[tag]
            if data.get("unit") and data["unit"] != l.unit:
                issues.append(
                    f"{tag}: changing price or quantity units requires a reviewed conversion."
                )
                continue
            for field, typ in [
                ("quantity", "set_quantity"),
                ("description", "update_description"),
                ("lead_time", "set_lead_time"),
            ]:
                if not data.get(field):
                    continue
                old = getattr(l, field)
                new = data[field]
                from backend.domain.pricing import dec

                if dec(new) != old if field == "quantity" else new != old:
                    operations.append(
                        Operation(
                            type=typ,
                            line_id=l.id,
                            value=new,
                            expected_before=str(old),
                            source_ids=[source],
                            reason=f"Reviewed mapping from {d.name}",
                        ).model_dump(mode="json")
                    )
        else:
            inp = LineInput(
                tag=tag,
                description=data.get("description")
                or f"{tag} — description needs review",
                quantity=data["quantity"],
                unit=data.get("unit") or "each",
                model=data.get("model", ""),
                cost=data.get("cost") or None,
                price=data.get("price") or None,
                lead_time=data.get("lead_time") or "Needs confirmation",
                source_ids=[source],
            )
            operations.append(
                Operation(type="add_line", line=inp, source_ids=[source]).model_dump(
                    mode="json"
                )
            )
    if issues:
        fail(
            {
                "message": "Resolve these source rows before continuing. No mapping was applied.",
                "issues": issues,
            }
        )
    d.meta = {**d.meta, "confirmed_mapping": mapping, "purpose": purpose}
    d.state = "mapped"
    if operations:
        operations.append(
            Operation(
                type="reconcile_inputs",
                reason=f"Accepted explicit changes from {d.name}; omitted rows are preserved.",
            ).model_dump(mode="json")
        )
        pr = Proposal(
            quote_id=q.id,
            base_version=q.version,
            input_revision=p.input_revision,
            title="Requote from " + d.name
            if purpose == "addendum"
            else "Draft from " + d.name,
            summary=f"{len(operations) - 1} sourced changes. Confirm scope and document precedence before accepting. Unlisted equipment is preserved.",
            operations=operations,
            sources=[r["source_id"] for r in rows],
        )
        s.add(pr)
        s.flush()
        return pr.id
    return None
