"""Deliberately small output projections from an order workspace.

Only accepted responses and engineering configuration deltas are exported.
Agent drafts, internal pricing and commercial calculations are not included.
"""

from __future__ import annotations

from html import escape
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from backend.domain.exports import safe_cell
from backend.orders.extraction import ATTRIBUTE_REGISTRY


def _display(value: Any) -> str:
    if value is None:
        return "Not stated"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _reference(workspace, source_ids):
    sources = {source["id"]: source for source in workspace.get("sources", [])}
    documents = {
        document["id"]: document for document in workspace.get("documents", [])
    }
    references = []
    for source_id in source_ids or []:
        source = sources.get(source_id)
        if not source:
            references.append(f"Evidence {source_id}")
            continue
        location = source.get("location") or {}
        document = documents.get(source.get("document_id"), {})
        where = (
            f"page {location['page']}"
            if location.get("page")
            else f"row {location['row']}"
            if location.get("row")
            else f"line {location['line']}"
            if location.get("line")
            else "source extract"
        )
        annotation = (
            f", annotation {location['annotation_id']}"
            if location.get("annotation_id")
            else ""
        )
        references.append(
            f"{document.get('name', 'Source document')} · {where}{annotation}"
        )
    return "; ".join(dict.fromkeys(references))


def response_matrix_rows(workspace: dict) -> list[dict]:
    # Response approval and drawing compliance are separate states. Use the
    # deterministic comment check so a passing value is visible even while its
    # outward response awaits approval, and an accepted response cannot hide a
    # drawing that still fails.
    checks = {check["id"]: check for check in workspace.get("checks", [])}
    labels = {
        "pass": "Verified",
        "fail": "Not addressed",
        "conflict": "Decision required",
        "unknown": "Needs review",
        "waived": "Signed off",
    }
    rows = []
    for comment in workspace.get("comments", []):
        check = checks.get("review:" + str(comment.get("id", "")), {})
        status = labels.get(
            check.get("status"), comment.get("status", "open").replace("_", " ")
        )
        ids = list(
            dict.fromkeys(
                [
                    *comment.get("source_ids", []),
                    *check.get("source_ids", []),
                    *check.get("actual_source_ids", []),
                ]
            )
        )
        rows.append(
            {
                "number": comment.get("number", ""),
                "device": comment.get("device") or "Needs identification",
                "comment": comment.get("text", ""),
                "status": status,
                "response": comment.get("response") or "",
                "evidence": _reference(workspace, ids),
            }
        )
    return rows


def response_matrix_html(workspace: dict) -> str:
    order = workspace.get("order") or {}
    e = lambda value: escape(str(value))
    rows = "".join(
        f"<tr><td class='number'>{e(row['number'])}</td><td><b>{e(row['device'])}</b><p>{e(row['comment'])}</p><small>{e(row['evidence'])}</small></td><td><span class='status'>{e(row['status'])}</span></td><td>{e(row['response']) if row['response'] else '<span class="empty">Awaiting an accepted response</span>'}</td></tr>"
        for row in response_matrix_rows(workspace)
    )
    demo = (
        "<aside>FICTIONAL REPLAY · Demonstration evidence only. These ratings are not a design certification or commercial offer.</aside>"
        if order.get("synthetic")
        else ""
    )
    incomplete = any(not row["response"] for row in response_matrix_rows(workspace))
    state = (
        "DRAFT · Responses remain to be accepted"
        if incomplete
        else "Accepted responses · Review before transmittal"
    )
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
    @page{{size:A4 landscape;margin:16mm}}*{{box-sizing:border-box}}body{{font:10px Arial,sans-serif;color:#20202a}}header{{display:flex;justify-content:space-between;border-bottom:2px solid #252334;padding-bottom:16px}}.brand{{font-size:26px;font-weight:bold;letter-spacing:-1px}}.eyebrow{{font-size:9px;letter-spacing:1px;text-transform:uppercase;color:#605b6e}}h1{{font-size:26px;letter-spacing:-.7px;margin:23px 0 8px}}.meta{{display:flex;gap:26px;color:#555261;margin:0 0 18px}}aside{{padding:10px 13px;background:#f2eefb;border-left:3px solid #7761a9;margin-bottom:18px}}table{{border-collapse:collapse;table-layout:fixed;width:100%}}thead{{display:table-header-group}}th{{padding:10px;text-align:left;font-size:9px;background:#282433;color:white}}th:nth-child(1){{width:5%}}th:nth-child(2){{width:45%}}th:nth-child(3){{width:13%}}th:nth-child(4){{width:37%}}td{{vertical-align:top;border-bottom:1px solid #dedbe5;padding:10px 10px;line-height:1.45;overflow-wrap:anywhere;white-space:pre-wrap}}tr{{break-inside:avoid}}p{{margin:6px 0}}small{{display:block;color:#6f6a77;font-size:8px;white-space:normal}}.number{{font-weight:bold;color:#6d558f}}.empty{{color:#797382;font-style:italic}}.status{{text-transform:capitalize}}footer{{margin-top:12px;color:#75717c;font-size:9px;border-top:1px solid #dedbe5;padding-top:12px}}
    </style></head><body><header><div class="brand">rivet.</div><div class="eyebrow">Comment response matrix<br>{e(order.get("number", "Order"))} · {e(order.get("revision_label") or "Revision unspecified")}</div></header><h1>{e(order.get("title", "Order response matrix"))}</h1><div class="meta"><span>{e(order.get("customer", ""))}</span><span>{e(state)}</span><span>Order state {e(order.get("version", "—"))}</span></div>{demo}<table><thead><tr><th>#</th><th>Review comment / evidence</th><th>Check result</th><th>Accepted response</th></tr></thead><tbody>{rows or '<tr><td colspan="4">No review comments in this order revision.</td></tr>'}</tbody></table><footer>Generated from the recorded order state. A response does not establish drawing compliance; closure is evaluated against revision evidence. This document does not authorize production release.</footer></body></html>"""


def response_matrix_pdf(workspace: dict) -> bytes:
    from pathlib import Path

    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        arguments = {"headless": True}
        chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
        if chrome.exists():
            arguments["executable_path"] = str(chrome)
        browser = playwright.chromium.launch(**arguments)
        try:
            page = browser.new_page()
            page.set_content(response_matrix_html(workspace), wait_until="load")
            return page.pdf(
                format="A4",
                landscape=True,
                print_background=True,
                display_header_footer=True,
                header_template="<span></span>",
                footer_template='<div style="font-size:8px;width:100%;text-align:center;color:#888">Rivet · <span class="pageNumber"></span> / <span class="totalPages"></span></div>',
            )
        finally:
            browser.close()


def _sheet(workspace, title, headings):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    order = workspace.get("order") or {}
    sheet.append(["RIVET | " + title.upper()])
    sheet.append(
        [safe_cell(order.get("number", "Order")), safe_cell(order.get("title", ""))]
    )
    sheet.append(["Customer", safe_cell(order.get("customer", ""))])
    sheet.append(
        [
            "Revision",
            safe_cell(order.get("revision_label") or "Unspecified"),
            "Order state",
            order.get("version", ""),
        ]
    )
    if order.get("synthetic"):
        sheet.append(
            ["FICTIONAL REPLAY — example ratings are not design certification"]
        )
    sheet.append([])
    sheet.append(headings)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="282433")
    sheet.freeze_panes = f"A{header_row + 1}"
    return workbook, sheet, header_row


def _save(workbook, sheet, widths, header_row):
    for number, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter

        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.auto_filter.ref = (
        f"A{header_row}:{sheet.cell(sheet.max_row, len(widths)).coordinate}"
    )
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.sheet_view.showGridLines = False
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def response_matrix_xlsx(workspace: dict) -> bytes:
    workbook, sheet, header = _sheet(
        workspace,
        "Response matrix",
        [
            "Comment",
            "Device",
            "Review comment",
            "Check result",
            "Accepted response",
            "Evidence",
        ],
    )
    for row in response_matrix_rows(workspace):
        sheet.append(
            [
                safe_cell(row[key])
                for key in (
                    "number",
                    "device",
                    "comment",
                    "status",
                    "response",
                    "evidence",
                )
            ]
        )
    return _save(workbook, sheet, [12, 20, 65, 20, 65, 65], header)


def bom_delta_xlsx(workspace: dict) -> bytes:
    workbook, sheet, header = _sheet(
        workspace,
        "BOM delta",
        ["Device", "Attribute", "Previous", "Current", "Unit", "Current evidence"],
    )
    changes = (workspace.get("diff") or {}).get("changes", [])
    for change in changes:
        attribute = change.get("attribute")
        # Never serialize a raw graph/change object. The allowlist intentionally
        # excludes cost, price, margin, agent reasoning and supplier economics.
        if attribute not in ATTRIBUTE_REGISTRY or change.get("before") == change.get(
            "after"
        ):
            # A new citation alone is a revision diff, but not a BOM change.
            continue
        definition = ATTRIBUTE_REGISTRY[attribute]
        sheet.append(
            [
                safe_cell(change.get("device", "")),
                safe_cell(definition["label"]),
                safe_cell(_display(change.get("before"))),
                safe_cell(_display(change.get("after"))),
                safe_cell(change.get("unit") or definition["unit"] or ""),
                safe_cell(_reference(workspace, change.get("source_ids", []))),
            ]
        )
    if sheet.max_row == header:
        sheet.append(["No configuration changes between the selected revisions."])
    return _save(workbook, sheet, [22, 35, 35, 35, 12, 75], header)
