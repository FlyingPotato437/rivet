from decimal import Decimal
from io import BytesIO
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from backend.domain.pricing import money

CUSTOMER_FIELDS = (
    "tag",
    "description",
    "model",
    "quantity",
    "unit",
    "price",
    "extended",
    "lead_time",
)


def customer_snapshot(snapshot):
    return {
        **{
            k: snapshot[k]
            for k in (
                "number",
                "version",
                "input_revision",
                "project",
                "customer",
                "currency",
                "terms",
                "synthetic",
            )
        },
        "lines": [{k: l[k] for k in CUSTOMER_FIELDS} for l in snapshot["lines"]],
    }


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def xlsx(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws.append(["RIVET | EQUIPMENT QUOTE"])
    ws.append([data["number"], f"Revision {data['version']}"])
    if data["synthetic"]:
        ws.append(["SYNTHETIC DEMONSTRATION — fictional equipment and pricing"])
    ws.append(["Project", safe_cell(data["project"])])
    ws.append(["Prepared for", safe_cell(data["customer"])])
    ws.append([])
    ws.append(
        [
            "Tag",
            "Description",
            "Model / configuration",
            "Quantity",
            "Unit",
            "Unit price (USD)",
            "Amount (USD)",
            "Delivery wording",
        ]
    )
    header = ws.max_row
    for l in data["lines"]:
        ws.append(
            [
                safe_cell(l["tag"]),
                safe_cell(l["description"]),
                safe_cell(l["model"]),
                Decimal(l["quantity"]),
                safe_cell(l["unit"]),
                Decimal(l["price"]) if l["price"] is not None else None,
                Decimal(l["extended"]) if l["extended"] is not None else None,
                safe_cell(l["lead_time"]),
            ]
        )
        for c in (6, 7):
            ws.cell(ws.max_row, c).number_format = '"$"#,##0.00'
    ws.append(
        [
            "",
            "",
            "",
            "",
            "",
            "TOTAL (USD)",
            sum(Decimal(l["extended"] or "0") for l in data["lines"]),
        ]
    )
    ws.cell(ws.max_row, 7).number_format = '"$"#,##0.00'
    ws.append([])
    ws.append(["Terms and exclusions", safe_cell(data["terms"])])
    for cell in ws[header]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17151E")
    for col, width in zip("ABCDEFGH", [16, 48, 25, 12, 12, 20, 20, 45]):
        ws.column_dimensions[col].width = width
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = f"A{header + 1}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def html(data):
    e = lambda v: escape(str(v))
    amount = lambda v: f"${Decimal(v or '0'):,.2f}"
    rows = "".join(
        f"<tr><td><b>{e(l['tag'])}</b></td><td>{e(l['description'])}<small>{e(l['model'])}</small></td><td>{e(l['quantity'])} {e(l['unit'])}</td><td>{amount(l['price'])}</td><td>{amount(l['extended'])}</td></tr>"
        for l in data["lines"]
    )
    delivery = "".join(
        f"<li><b>{e(l['tag'])}</b> — {e(l['lead_time'])}</li>" for l in data["lines"]
    )
    total = amount(sum(Decimal(l["extended"] or "0") for l in data["lines"]))
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>@page{{size:A4;margin:18mm}}body{{font-family:Arial,sans-serif;color:#16151c;font-size:10px}}header{{display:flex;justify-content:space-between;border-bottom:2px solid;padding-bottom:18px}}h1{{font-size:30px;letter-spacing:-1px;margin:0}}h2{{font-size:23px;margin:30px 0 8px}}small{{display:block;color:#666;margin-top:5px}}table{{width:100%;border-collapse:collapse;margin:25px 0}}th{{text-align:left;background:#eeedf1;padding:10px 7px;font-size:9px;text-transform:uppercase}}td{{border-bottom:1px solid #ddd;padding:12px 7px;vertical-align:top}}tr{{break-inside:avoid}}thead{{display:table-header-group}}.total{{text-align:right;font-size:18px}}.demo{{padding:10px;background:#f3edf9;margin:18px 0}}p,li{{line-height:1.7}}footer{{margin-top:30px;border-top:1px solid #ddd;padding-top:10px;color:#777}}</style></head><body><header><h1>rivet.</h1><div>EQUIPMENT QUOTATION<br>{e(data["number"])} / REV {data["version"]:02d}</div></header>{'<div class="demo">SYNTHETIC DEMONSTRATION · Fictional equipment, prices and approvals. Not a commercial offer.</div>' if data["synthetic"] else ""}<h2>{e(data["project"])}</h2><p>Prepared for {e(data["customer"])}</p><table><thead><tr><th>Tag</th><th>Equipment / configuration</th><th>Quantity</th><th>Unit price</th><th>Amount</th></tr></thead><tbody>{rows}</tbody></table><p class="total">Total USD &nbsp; <b>{total}</b></p><h3>Delivery</h3><ul>{delivery}</ul><h3>Terms & exclusions</h3><p>{e(data["terms"])}</p><footer>Rivet · Approved revision {data["version"]} · Input revision {data["input_revision"]} · Generated from an immutable quote snapshot</footer></body></html>"""


def pdf(data):
    from playwright.sync_api import sync_playwright
    import os

    with sync_playwright() as p:
        args = {"headless": True}
        chrome = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
        if os.path.exists(chrome):
            args["executable_path"] = chrome
        browser = p.chromium.launch(**args)
        page = browser.new_page()
        page.set_content(html(data), wait_until="load")
        result = page.pdf(
            format="A4",
            print_background=True,
            display_header_footer=True,
            header_template="<span></span>",
            footer_template='<div style="font-size:8px;width:100%;text-align:center;color:#888"><span class="pageNumber"></span> / <span class="totalPages"></span></div>',
        )
        browser.close()
        return result
