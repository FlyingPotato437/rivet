import re
from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from backend.domain.exports import safe_cell

HEADINGS = [
    "Comment",
    "Revision",
    "Source file",
    "Source page",
    "Drawing location",
    "Comment text",
    "Original text",
    "Author",
    "Date",
    "Status",
    "Response",
    "Response author",
    "Review",
    "Linked email",
]


def rows(w, compact=False):
    docs = {d["id"]: d for d in w["documents"]}
    refs = {id: f"D{i}" for i, id in enumerate(docs, 1)}
    sources = {sp["id"]: sp for sp in w["sources"]}
    order = {id: i for i, id in enumerate(docs)}
    for c in sorted(
        w["comments"],
        key=lambda c: (
            order.get(c["document_id"], -1),
            tuple(
                int(p) if p.isdigit() else p.lower()
                for p in re.split(r"(\d+)", c["number"])
            ),
        ),
    ):
        target = sources.get(c["target_source_id"], {})
        location = target.get("location", {})
        target_label = (
            (
                (
                    refs.get(target.get("document_id"), "")
                    if compact
                    else docs.get(target.get("document_id"), {}).get("name", "")
                )
                + " · "
                + (
                    f"page {location['page']}"
                    if location.get("page")
                    else location.get("label", "source excerpt")
                )
            )
            if target
            else "Not linked"
        )
        yield [
            c["number"],
            c["revision"],
            refs.get(c["document_id"], "Manual entry")
            if compact
            else docs.get(c["document_id"], {}).get("name", "Manual entry"),
            c["source_page"] or "Not stated",
            target_label,
            c["text"],
            c["original_text"],
            c["author"] or "Not stated",
            c["authored_at"] or "Not stated",
            c["status"],
            c["response"],
            c["responder"],
            c["confidence"],
            docs.get(c["linked_email_id"], {}).get("name", ""),
        ]


def workbook(w):
    book = Workbook()
    sheet = book.active
    sheet.title = "Comment log"
    sheet.append([safe_cell(w["order"]["title"]), safe_cell(w["order"]["number"])])
    sheet.append(HEADINGS)
    for row in rows(w):
        sheet.append([safe_cell(v) if isinstance(v, str) else v for v in row])
    sheet.freeze_panes = "F3"
    sheet.auto_filter.ref = f"A2:N{sheet.max_row}"
    for i, width in enumerate(
        [12, 15, 32, 14, 40, 65, 65, 25, 28, 15, 65, 25, 20, 30], 1
    ):
        sheet.column_dimensions[sheet.cell(2, i).column_letter].width = width
    changes = book.create_sheet("Change record")
    changes.append(
        [
            "Change",
            "Before",
            "After",
            "From revision",
            "To revision",
            "Linked comments",
            "Requested by",
            "Approved by",
            "Approved at",
        ]
    )
    byid = {c["id"]: c["number"] for c in w["comments"]}
    for c in w["changes"]:
        changes.append(
            [
                safe_cell(v)
                for v in [
                    c["title"],
                    c["before"],
                    c["after"],
                    c["from_revision"],
                    c["to_revision"],
                    ", ".join(byid[i] for i in c["comment_ids"] if i in byid),
                    c["requested_by"],
                    c["approved_by"],
                    c["approved_at"],
                ]
            ]
        )
    for sh in book:
        for row in sh:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == (2 if sh == sheet else 1):
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="34323F")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def pdf(w):
    from pathlib import Path

    from playwright.sync_api import sync_playwright

    e = lambda value: escape(str(value))
    source_key = "".join(
        f"<p><b>D{i}</b> {e(d['name'])}</p>" for i, d in enumerate(w["documents"], 1)
    )
    body = "".join(
        f"<tr><td>{e(r[0])}</td><td>{e(r[2])} · Page {e(r[3])}<br>Revision: {e(r[1])}<br>Target: {e(r[4])}</td><td>{e(r[5])}<p>{e(r[7])} · {e(r[8])}</p></td><td>{e(r[9])}<br>{e(r[12])}</td><td>{e(r[10])}<p>{e(r[11])}</p></td></tr>"
        for r in rows(w, compact=True)
    )
    html = f"""<html><head><title>{e(w["order"]["title"])} - comment log</title><style>@page{{size:A4 landscape;margin:16mm}}body{{font:11px/1.45 Arial;color:#25232b}}h1{{font-size:24px}}table{{border-collapse:collapse;width:100%;table-layout:fixed}}td,th{{border-bottom:1px solid #ddd;padding:8px;vertical-align:top;text-align:left;overflow-wrap:anywhere;white-space:pre-wrap}}th{{background:#eee}}tr{{break-inside:avoid}}p{{color:#5c5864}}</style></head><body><h1>{e(w["order"]["title"])}</h1><p>{e(w["order"]["number"])} · Comment response matrix · {e(w["order"]["customer"])}</p><table><colgroup><col style="width:7%"><col style="width:19%"><col style="width:32%"><col style="width:12%"><col style="width:30%"></colgroup><thead><tr><th>No.</th><th>Source and location</th><th>Comment / author / date</th><th>Status</th><th>Response / author</th></tr></thead><tbody>{body}</tbody></table><h3>Source files</h3>{source_key}<p>Recorded responses and review status. This record does not certify equipment compliance or authorize production.</p></body></html>"""
    with sync_playwright() as pw:
        chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
        browser = pw.chromium.launch(
            **({"executable_path": str(chrome)} if chrome.exists() else {})
        )
        try:
            page = browser.new_page()
            page.set_content(html)
            return page.pdf(format="A4", landscape=True, print_background=True)
        finally:
            browser.close()
