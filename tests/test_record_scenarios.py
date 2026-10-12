"""Document-to-record scenarios using generated files and the real import worker.

Fixtures are explicitly fictional. No model calls, outbound messages, or public
document downloads are needed for these integration checks.
"""

import re
from email.message import EmailMessage
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from test_records import create, edit_body, post, view

from backend.domain.service import get
from backend.records.email_intake import import_email
from backend.storage.db import Session
from backend.storage.models import Document, Job, Order
from backend.worker import claim, process
from scripts.make_order_replay import write_pdf


def rewrite_pdf_objects(path, transform):
    """Edit only our simple generated fixtures, rebuilding a valid xref table."""
    objects = re.findall(rb"\d+ 0 obj\n(.*?)\nendobj\n", path.read_bytes(), re.DOTALL)
    transform(objects)
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    )
    path.write_bytes(result)


def drain():
    while task := claim():
        process(task)


def import_file(client, order, name, content, *, role="markups", revision="B"):
    response = client.post(
        f"/api/projects/{order['project_id']}/documents",
        files={"file": (name, content, "application/octet-stream")},
        data={"kind": role, "revision_label": revision},
        headers={"Idempotency-Key": str(uuid4())},
    )
    assert response.status_code == 202, response.text
    drain()
    return response.json(), view(client, order)


def test_hundred_page_package_keeps_markup_page_separate_from_target_drawing(
    client, tmp_path
):
    # The generator places 35 entries per page. Blank entries keep this a small
    # native PDF while exercising all 100 physical pages, including the last.
    lines = []
    for page in range(1, 101):
        lines.extend([f"Drawing E-{page:02d} | Revision B", *([""] * 34)])
    path = tmp_path / "fictional-100-page-submittal.pdf"
    text = "Please confirm cable entry on drawing E-10, page 10."
    write_pdf(
        path,
        "Fictional custom switchboard submittal",
        lines,
        [
            {
                "id": "review-7",
                "page": 2,
                "text": "Comment 7: " + text,
                "author": "Casey Review",
                "date": "D:20261002120000Z",
            }
        ],
    )
    order = create(client)
    _, workspace = import_file(
        client, order, path.name, path.read_bytes(), role="approval_drawing"
    )
    document = workspace["documents"][0]
    processed = [c for c in document["coverage"] if c["state"] == "processed"]
    assert document["state"] == "ready" and len(processed) == 100
    assert processed[-1]["label"] == "Page 100"
    assert len(workspace["comments"]) == 1
    comment = workspace["comments"][0]
    source = next(s for s in workspace["sources"] if s["id"] in comment["source_ids"])
    assert comment["source_page"] == source["location"]["page"] == 2
    assert source["location"]["annotation_id"] == "review-7"
    assert not comment["reviewed"]
    assert "page 10" in " ".join(comment["flags"])
    drawing = next(
        s
        for s in workspace["sources"]
        if s["location"].get("page") == 10 and "Drawing E-10" in s["text"]
    )
    # An exact identifier can be linked by the deterministic coordinator, but
    # remains visibly automatic and unreviewed; the markup location never moves.
    assert comment["target_source_id"] == drawing["id"]
    assert comment["target_source_id"] not in comment["source_ids"]
    assert any(link["automatic"] for link in workspace["coordination"]["links"])
    response = post(
        client,
        f"/orders/{order['id']}/record/comments/{comment['id']}",
        edit_body(workspace, comment, target_source_id=drawing["id"], reviewed=True),
    )
    assert response.status_code == 200, response.text
    saved = response.json()["comments"][0]
    assert saved["source_page"] == 2 and saved["target_source_id"] == drawing["id"]
    assert saved["original_text"] == text and saved["status"] == "open"
    exported = client.get(f"/api/orders/{order['id']}/record/export?format=xlsx")
    assert exported.status_code == 200
    row = list(load_workbook(BytesIO(exported.content)).active.values)[2]
    assert row[3] == 2 and "page 10" in row[4]


def test_same_comment_number_in_new_revision_keeps_prior_response_separate(
    client, tmp_path
):
    order = create(client)
    workspaces = []
    for revision in ("B", "C"):
        path = tmp_path / f"returned-Rev-{revision}.pdf"
        write_pdf(
            path,
            "Fictional returned review",
            [f"Drawing E-10 | Revision {revision}"],
            [{"id": "comment-7", "text": "Comment 7: Confirm cable entry."}],
        )
        _, workspace = import_file(
            client, order, path.name, path.read_bytes(), revision=revision
        )
        if revision == "B":
            original = workspace["comments"][0]
            result = post(
                client,
                f"/orders/{order['id']}/record/comments/{original['id']}",
                edit_body(
                    workspace,
                    original,
                    reviewed=True,
                    status="responded",
                    response="Rev B response recorded; engineering confirmation is pending.",
                    responder="Casey Review",
                ),
            )
            assert result.status_code == 200, result.text
            workspace = result.json()
        workspaces.append(workspace)
    comments = workspaces[-1]["comments"]
    assert len(comments) == 2 and {c["number"] for c in comments} == {"7"}
    assert len({c["id"] for c in comments}) == 2
    prior = next(c for c in comments if c["revision"] == "B")
    newer = next(c for c in comments if c["revision"] == "C")
    assert prior["response"] == workspaces[0]["comments"][0]["response"]
    assert prior["status"] == "responded" and prior["reviewed"]
    assert newer["status"] == "open" and not newer["reviewed"]
    assert newer["response"] == newer["target_source_id"] == ""
    assert prior["source_ids"] != newer["source_ids"]


def test_exact_duplicate_upload_is_noop_but_distinct_revision_and_order_are_retained(
    client,
):
    order = create(client)
    content = b"Comment 7: Confirm cable entry.\n"
    first, original = import_file(client, order, "review.txt", content)
    with Session() as s:
        jobs_before = s.scalar(select(func.count()).select_from(Job))
    duplicate, repeated = import_file(client, order, "review.txt", content)
    assert duplicate["id"] == first["id"] and duplicate["duplicate"] is True
    assert repeated == original
    with Session() as s:
        assert s.scalar(select(func.count()).select_from(Job)) == jobs_before
    revision, revised = import_file(client, order, "review.txt", content, revision="C")
    assert revision["id"] != first["id"] and len(revised["comments"]) == 2
    other = create(client)
    separate, independent = import_file(client, other, "review.txt", content)
    assert separate["id"] != first["id"] and len(independent["comments"]) == 1


def test_unstructured_email_stays_one_unresolved_record_without_inferred_approval(
    client,
):
    order = create(client)
    body = (
        "The latest looks fine except the red item on the second sheet.\n"
        "Can you confirm this is the same as last time?\n\nRegards, Casey"
    )
    message = EmailMessage()
    message["From"] = "Casey <casey@example.com>"
    message["To"] = "pm@example.com"
    message["Date"] = "Fri, 02 Oct 2026 12:00:00 -0700"
    message["Subject"] = "Re: custom order"
    message.set_content(body)
    response = client.post(
        f"/api/orders/{order['id']}/record/email",
        files={"file": ("ambiguous.eml", message.as_bytes(), "message/rfc822")},
        headers={"Idempotency-Key": str(uuid4())},
    )
    assert response.status_code == 202, response.text
    drain()
    workspace = view(client, order)
    assert len(workspace["comments"]) == 1
    comment = workspace["comments"][0]
    assert comment["text"] == comment["original_text"] == body
    assert comment["origin"] == "email" and comment["linked_email_id"]
    assert comment["source_page"] is None and comment["target_source_id"] == ""
    assert not comment["reviewed"] and comment["status"] == "open"
    assert not comment["response"] and not workspace["approvals"]


def test_spreadsheets_keep_row_evidence_and_uncached_formula_requires_review(client):
    order = create(client)
    _, workspace = import_file(
        client,
        order,
        "equipment.csv",
        b"tag,description,quantity\nMSB-01,Main switchboard,2\n",
        role="bom",
    )
    csv_source = next(s for s in workspace["sources"] if "MSB-01" in s["text"])
    assert csv_source["location"]["row"] == 2
    assert csv_source["location"]["sheet"] == "CSV"
    assert "quantity: 2" in csv_source["text"]
    assert not workspace["comments"]  # Equipment rows are not reviewer comments.
    book = Workbook()
    book.active.title = "Equipment"
    book.active.append(["tag", "quantity"])
    book.active.append(["MSB-01", "=1+1"])
    content = BytesIO()
    book.save(content)
    _, workspace = import_file(
        client, order, "equipment.xlsx", content.getvalue(), role="bom"
    )
    document = next(d for d in workspace["documents"] if d["name"] == "equipment.xlsx")
    assert any(c["state"] == "unsupported" for c in document["coverage"])
    source = next(s for s in workspace["sources"] if s["document_id"] == document["id"])
    assert source["location"]["formulas"] == {"B2": "=1+1"}
    assert "quantity: 2" not in source["text"]
    review = next(
        c for c in workspace["comments"] if c["document_id"] == document["id"]
    )
    assert review["origin"] == "unreadable area" and not review["reviewed"]
    assert not workspace["approvals"]


def test_forwarded_unsupported_attachment_is_visible_and_parent_provenance_is_distinct(
    client,
):
    order = create(client)
    parents = []
    for index in (1, 2):
        message = EmailMessage()
        message["From"] = "casey@example.com"
        message["To"] = "pm@example.com"
        message["Subject"] = f"Review delivery {index}"
        message["Message-ID"] = f"<review-{index}@example.com>"
        message.set_content("Please see the attached review files.")
        message.add_attachment(
            b"Comment 7: Confirm cable entry.",
            maintype="text",
            subtype="plain",
            filename="review.txt",
        )
        message.add_attachment(
            b"fictional unsupported CAD content",
            maintype="application",
            subtype="octet-stream",
            filename="mechanical-layout.dwg",
        )
        with Session.begin() as s:
            result = import_email(
                s,
                get(s, Order, order["id"]),
                f"forward-{index}.eml",
                message.as_bytes(),
                forwarded=True,
            )
            parents.append(result["document_id"])
            assert result["unparsed_attachments"] == ["mechanical-layout.dwg"]
        drain()
    workspace = view(client, order)
    emails = [d for d in workspace["documents"] if d["id"] in parents]
    assert len(emails) == 2
    for document in emails:
        assert any(
            "mechanical-layout.dwg" in warning for warning in document["warnings"]
        )
        assert all("CAD content" not in warning for warning in document["warnings"])
    attachments = [d for d in workspace["documents"] if d["name"] == "review.txt"]
    assert len(attachments) == 2
    with Session() as s:
        assert {
            get(s, Document, d["id"]).meta["email_document_id"] for d in attachments
        } == set(parents)
    assert not workspace["approvals"]


def test_pdf_popup_and_form_widgets_are_not_extra_review_comments(client, tmp_path):
    path = tmp_path / "review-with-form.pdf"
    write_pdf(
        path,
        "Fictional marked-up form",
        ["Drawing E-10"],
        [
            {"id": "parent", "text": "Comment 7: Confirm cable entry."},
            {"id": "popup", "kind": "Popup", "text": ""},
            {"id": "form-field", "kind": "Widget", "text": ""},
            {"id": "website", "kind": "Link", "text": ""},
        ],
    )
    order = create(client)
    _, workspace = import_file(client, order, path.name, path.read_bytes())
    assert len(workspace["comments"]) == 1
    assert workspace["comments"][0]["number"] == "7"
    annotations = [
        s for s in workspace["sources"] if s["location"].get("annotation_id")
    ]
    assert [s["location"]["annotation_id"] for s in annotations] == ["parent"]


def test_flattened_unnumbered_markup_reports_limit_without_inventing_comments(
    client, tmp_path
):
    path = tmp_path / "flattened-markup.pdf"
    write_pdf(
        path, "Fictional flattened markup", ["Please confirm the red cable entry."]
    )
    order = create(client)
    _, workspace = import_file(client, order, path.name, path.read_bytes())
    assert not workspace["comments"]
    document = workspace["documents"][0]
    assert document["state"] == "ready"
    assert any(
        "no comments were inferred" in warning for warning in document["warnings"]
    )
    assert any(
        c["label"] == "Markup extraction" and c["state"] == "unsupported"
        for c in document["coverage"]
    )
    assert workspace["sources"] and not workspace["approvals"]


def test_native_and_scanned_pages_in_one_package_preserve_page_specific_coverage(
    client, tmp_path
):
    path = tmp_path / "mixed-package.pdf"
    write_pdf(
        path, "Fictional mixed package", ["Comment 1: Confirm the panel.", *([""] * 69)]
    )

    def add_scan(objects):
        pages = [i for i, obj in enumerate(objects) if b"/Type /Page " in obj]
        page_index = pages[1]
        content_id = int(re.search(rb"/Contents (\d+) 0 R", objects[page_index])[1])
        image_id = len(objects) + 1
        pixels = b"\x22\x22\x22\xff\xff\xff\xff\xff\xff\x22\x22\x22"
        objects.append(
            b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 12 >>\nstream\n"
            + pixels
            + b"\nendstream"
        )
        objects[page_index] = objects[page_index].replace(
            b"/Resources << /Font << /F1 3 0 R >> >>",
            f"/Resources << /XObject << /Scan {image_id} 0 R >> >>".encode(),
        )
        content = b"q 450 0 0 650 70 80 cm /Scan Do Q"
        objects[content_id - 1] = (
            f"<< /Length {len(content)} >>\nstream\n".encode()
            + content
            + b"\nendstream"
        )

    rewrite_pdf_objects(path, add_scan)
    order = create(client)
    _, workspace = import_file(client, order, path.name, path.read_bytes())
    coverage = workspace["documents"][0]["coverage"]
    assert any(c["label"] == "Page 1" and c["state"] == "processed" for c in coverage)
    assert any(c["label"] == "Page 2" and c["state"] == "scanned" for c in coverage)
    scan = next(c for c in workspace["comments"] if c["origin"] == "unreadable area")
    assert scan["source_page"] == 2 and not scan["reviewed"]
    assert scan["original_text"] == "" and scan["source_ids"]


def geometry_fixture(path, rotation, nonzero_origin=True):
    """Known raw PDF coordinates; the pale square is the annotation target."""
    write_pdf(
        path,
        "Fictional PDF alignment QA",
        ["Drawing E-10"],
        [
            {
                "id": "alignment-box",
                "rect": [140, 220, 260, 340],
                "text": "Comment 1: Check the highlighted source area.",
            }
        ],
    )

    def transform(objects):
        index = next(i for i, obj in enumerate(objects) if b"/Type /Page " in obj)
        media = b"20 30 632 872" if nonzero_origin else b"0 0 612 842"
        objects[index] = objects[index].replace(
            b"/MediaBox [0 0 612 842]",
            b"/MediaBox ["
            + media
            + b"] /CropBox [70 80 550 720] /Rotate "
            + str(rotation).encode(),
        )
        content_id = int(re.search(rb"/Contents (\d+) 0 R", objects[index])[1])
        content = (
            b"q 0.87 0.84 0.96 rg 140 220 120 120 re f 0.3 0.2 0.5 RG 1 w 140 220 120 120 re S Q\n"
            b"BT /F1 12 Tf 140 300 Td (SOURCE) Tj ET\n"
            b"BT /F1 16 Tf 100 680 Td (FICTIONAL PDF ALIGNMENT QA) Tj ET\n"
            b"BT /F1 10 Tf 100 650 Td (Crop + rotation: highlight should fit the pale square.) Tj ET\n"
            b"BT /F1 10 Tf 25 820 Td (HIDDEN OUTSIDE CROP) Tj ET\n"
        )
        objects[content_id - 1] = (
            f"<< /Length {len(content)} >>\nstream\n".encode()
            + content
            + b"\nendstream"
        )

    rewrite_pdf_objects(path, transform)


@pytest.mark.parametrize("nonzero_origin", [False, True])
@pytest.mark.parametrize(
    "rotation,expected",
    [
        (0, [70 / 480, 380 / 640, 190 / 480, 500 / 640]),
        (90, [140 / 640, 70 / 480, 260 / 640, 190 / 480]),
        (180, [290 / 480, 140 / 640, 410 / 480, 260 / 640]),
        (270, [380 / 640, 290 / 480, 500 / 640, 410 / 480]),
    ],
)
def test_rotated_cropped_pdf_matches_visible_viewport_coordinates(
    client, tmp_path, rotation, expected, nonzero_origin
):
    path = tmp_path / f"alignment-{rotation}.pdf"
    geometry_fixture(path, rotation, nonzero_origin)
    order = create(client)
    _, workspace = import_file(client, order, path.name, path.read_bytes())
    annotation = next(
        s
        for s in workspace["sources"]
        if s["location"].get("annotation_id") == "alignment-box"
    )
    assert annotation["location"]["bbox"] == pytest.approx(expected, abs=0.000001)
    assert annotation["location"]["approximate"] is False
    assert annotation["location"]["coordinate_system"] == "visible-page-v1"
    assert all("HIDDEN" not in s["text"] for s in workspace["sources"])
    # SOURCE starts exactly at raw PDF x=140. Its left/top/right/bottom axis
    # after rotation must agree with the same crop origin as the annotation.
    text = next(s for s in workspace["sources"] if s["text"] == "SOURCE")
    axis = {0: 0, 90: 1, 180: 2, 270: 3}[rotation]
    expected_edge = 70 / 480 if rotation in (0, 90) else 410 / 480
    assert text["location"]["bbox"][axis] == pytest.approx(expected_edge, abs=0.000001)
