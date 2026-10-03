from copy import deepcopy
from uuid import uuid4

from backend.domain.service import get
from backend.storage.db import Session
from backend.storage.models import Document, Span


def post(client, path, data, key=None):
    return client.post(
        "/api/orders" + path,
        json=data,
        headers={"Idempotency-Key": key or str(uuid4())},
    )


def create(client, title="Switchgear replay"):
    response = post(
        client, "", {"title": title, "customer": "Test Engineering", "synthetic": True}
    )
    assert response.status_code == 200, response.text
    return response.json()


def evidence(w, name, role, values, state="parsed"):
    with Session.begin() as s:
        document = Document(
            project_id=w["order"]["project_id"],
            name=name,
            kind=role,
            blob="fixture.txt",
            sha256=uuid4().hex * 2,
            state=state,
            meta={"order_role": role},
            coverage=[{"label": "Page 1", "state": "complete"}],
        )
        s.add(document)
        s.flush()
        ids = []
        for idx, value in enumerate(values):
            text = (
                value if isinstance(value, str) else f"MSB-01 | {value[0]}: {value[1]}"
            )
            span = Span(
                document_id=document.id,
                text=text,
                location={"page": 1, "bbox": [0.1, idx * 0.03, 0.8, idx * 0.03 + 0.02]},
            )
            s.add(span)
            s.flush()
            ids.append(span.id)
        return document.id, ids


def refresh(client, w):
    current = client.get("/api/orders/" + w["order"]["id"]).json()
    response = post(
        client,
        "/" + w["order"]["id"] + "/refresh",
        {"expected_version": current["order"]["version"]},
    )
    assert response.status_code == 200, response.text
    return response.json()


def fixture(client, actual=65):
    w = create(client)
    evidence(
        w,
        "Project specification.txt",
        "specification",
        [("short_circuit_ka", 65), ("bus_material", "copper")],
    )
    evidence(
        w, "Customer purchase order.txt", "purchase_order", [("short_circuit_ka", 85)]
    )
    evidence(
        w,
        "Approval drawing Rev B.txt",
        "drawing",
        [("short_circuit_ka", actual), ("bus_material", "copper")],
    )
    return refresh(client, w)


def decision(client, w, action, accept=True, key=None):
    response = post(
        client,
        f"/{w['order']['id']}/actions/{action['id']}/{'accept' if accept else 'reject'}",
        {
            "expected_version": w["order"]["version"],
            "reason": "Reviewed source evidence and engineering scope.",
        },
        key,
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_precedence_citations_versions_and_alias_approval(client):
    w = fixture(client)
    rating = next(r for r in w["ledger"] if r["attribute"] == "short_circuit_ka")
    assert (
        rating["expected"],
        rating["actual"],
        rating["precedence"],
        rating["status"],
    ) == (85, 65, "purchase_order", "fail")
    assert rating["source_ids"] and rating["actual_source_ids"]
    baseline = w["order"]["version"]
    before = client.get(f"/api/orders/{w['order']['id']}/revisions/{baseline}").json()
    evidence(
        w,
        "Accepted exceptions.txt",
        "accepted_exception",
        [("bus_material", "aluminum"), ("short_circuit_ka", 50)],
    )
    w = refresh(client, w)
    assert (
        next(r for r in w["ledger"] if r["attribute"] == "bus_material")["expected"]
        == "copper"
    )
    exception = next(
        a
        for a in w["actions"]
        if a["status"] == "pending" and "accept_exception_document" in a["after"]
    )
    w = decision(client, w, exception)
    assert (
        next(r for r in w["ledger"] if r["attribute"] == "bus_material")["expected"]
        == "aluminum"
    )
    assert (
        next(r for r in w["ledger"] if r["attribute"] == "short_circuit_ka")["expected"]
        == 85
    )
    assert (
        client.get(f"/api/orders/{w['order']['id']}/revisions/{baseline}").json()
        == before
    )
    difference = client.get(
        f"/api/orders/{w['order']['id']}/diff?from_version={baseline}"
    ).json()
    assert any(
        x["attribute"] == "bus_material" and x["expected_after"] == "aluminum"
        for x in difference["changes"]
    )


def test_actions_do_not_fake_drawings_and_propagation_requires_evidence(client):
    w = fixture(client)
    action = next(
        a
        for a in w["actions"]
        if a["status"] == "pending"
        and a["after"].get("attribute") == "short_circuit_ka"
    )
    key = str(uuid4())
    original = deepcopy(w)
    w = decision(client, w, action, key=key)
    assert (
        next(r for r in w["ledger"] if r["attribute"] == "short_circuit_ka")["actual"]
        == 65
    )
    assert len(w["tasks"]) == 4 and all(t["status"] == "pending" for t in w["tasks"])
    again = decision(client, original, action, key=key)
    assert again["order"]["version"] == w["order"]["version"]
    task = w["tasks"][0]
    response = post(
        client,
        f"/{w['order']['id']}/tasks/{task['id']}/verify",
        {"expected_version": w["order"]["version"], "source_ids": []},
    )
    assert response.status_code == 422 and "marking a task done" in response.text
    evidence(
        w,
        "Approval drawing Rev C.txt",
        "drawing",
        [("short_circuit_ka", 85), ("bus_material", "copper")],
    )
    w = refresh(client, w)
    assert (
        next(t for t in w["tasks"] if t["target"] == "drawing")["status"] == "verified"
    )
    assert sum(t["status"] == "pending" for t in w["tasks"]) == 3
    for role in ("bom", "supplier_po", "nameplate"):
        evidence(w, role + ".txt", role, [("short_circuit_ka", 85)])
    w = refresh(client, w)
    assert all(t["status"] == "verified" for t in w["tasks"])
    assert w["release"]["ready"] is True


def test_signed_waiver_fingerprint_and_new_input_invalidate_release(client):
    w = fixture(client)
    check = next(c for c in w["checks"] if c["attribute"] == "short_circuit_ka")
    body = {
        "expected_version": w["order"]["version"],
        "check_id": check["id"],
        "fingerprint": check["fingerprint"],
        "signer": "Test Engineer",
        "reason": "Signed exception for this exact documented condition.",
    }
    response = post(client, f"/{w['order']['id']}/waivers", body)
    assert response.status_code == 200, response.text
    w = response.json()
    assert next(c for c in w["checks"] if c["id"] == check["id"])["status"] == "waived"
    action = next(a for a in w["actions"] if a["status"] == "pending")
    w = decision(client, w, action, accept=False)
    assert w["release"]["ready"]
    release_body = {
        "expected_version": w["order"]["version"],
        "snapshot_hash": w["order"]["snapshot_hash"],
        "signer": "Test Engineer",
        "reason": "All checks reviewed and exact package authorized.",
    }
    response = post(client, f"/{w['order']['id']}/release", release_body)
    assert response.status_code == 200, response.text
    assert (
        response.json()["release"]["approval"]["snapshot_hash"]
        == w["order"]["snapshot_hash"]
    )
    evidence(
        w,
        "Approval drawing Rev C.txt",
        "drawing",
        [("short_circuit_ka", 50), ("bus_material", "copper")],
    )
    w = refresh(client, w)
    assert not w["release"]["approval"]
    assert next(c for c in w["checks"] if c["id"] == check["id"])["status"] == "fail"
    response = post(
        client,
        f"/{w['order']['id']}/waivers",
        body | {"expected_version": w["order"]["version"]},
    )
    assert response.status_code == 409
    assert post(client, f"/{w['order']['id']}/release", release_body).status_code == 409


def test_empty_or_unreadable_coverage_cannot_be_waived_or_released(client):
    w = create(client)
    check = next(c for c in w["checks"] if c["id"] == "coverage")
    assert check["status"] == "unknown" and not w["release"]["ready"]
    response = post(
        client,
        f"/{w['order']['id']}/waivers",
        {
            "expected_version": w["order"]["version"],
            "check_id": check["id"],
            "fingerprint": check["fingerprint"],
            "signer": "Test Engineer",
            "reason": "I would like to skip unread source coverage.",
        },
    )
    assert response.status_code == 422
    evidence(w, "Scanned drawing.pdf", "drawing", [], state="failed")
    w = refresh(client, w)
    assert (
        not w["release"]["ready"]
        and "could not"
        in next(c for c in w["checks"] if c["id"] == "coverage")["detail"]
    )


def test_cross_order_evidence_and_actions_rejected(client):
    w = fixture(client)
    other = fixture(client)
    action = next(a for a in other["actions"] if a["status"] == "pending")
    response = post(
        client,
        f"/{w['order']['id']}/actions/{action['id']}/accept",
        {"expected_version": w["order"]["version"], "reason": "Reviewed this change."},
    )
    assert response.status_code == 403
    response = post(
        client,
        f"/{w['order']['id']}/actions",
        {
            "expected_version": w["order"]["version"],
            "type": "configuration",
            "title": "Unsafe cross-order evidence",
            "after": action["after"],
            "source_ids": action["source_ids"],
        },
    )
    assert response.status_code == 403


def test_stale_revision_and_conflicting_current_sources(client):
    w = fixture(client)
    old = w["order"]["version"]
    evidence(
        w,
        "Separate customer purchase order.txt",
        "purchase_order",
        [("short_circuit_ka", 100)],
    )
    w = refresh(client, w)
    assert (
        next(c for c in w["checks"] if c["attribute"] == "short_circuit_ka")["status"]
        == "conflict"
    )
    assert (
        post(
            client, f"/{w['order']['id']}/refresh", {"expected_version": old}
        ).status_code
        == 409
    )
    assert all(a["status"] == "stale" for a in w["actions"])


def test_new_drawing_does_not_inherit_omitted_old_values(client):
    w = fixture(client)
    evidence(w, "Approval drawing Rev C.txt", "drawing", [("bus_material", "copper")])
    w = refresh(client, w)
    rating = next(c for c in w["checks"] if c["attribute"] == "short_circuit_ka")
    assert rating["actual"] is None and rating["status"] == "unknown"


def test_comment_draft_acceptance_and_unverified_fixed_comment(client, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    w = fixture(client)
    evidence(
        w,
        "Reviewer markups.txt",
        "markup",
        [
            "Comment 7: MSB-01 | short_circuit_ka: 85",
            "Comment 8: Confirm lifting eyes.",
        ],
    )
    w = refresh(client, w)
    comment = next(c for c in w["comments"] if c["number"] == "7")
    response = post(
        client,
        f"/{w['order']['id']}/commands",
        {
            "expected_version": w["order"]["version"],
            "prompt": "Reply to comment 7: We will revise the drawing after engineering review.",
        },
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["actions"][0]["status"] == "pending"
    assert not next(
        c for c in result["workspace"]["comments"] if c["id"] == comment["id"]
    )["response"]
    w = decision(client, result["workspace"], result["actions"][0])
    assert next(c for c in w["comments"] if c["id"] == comment["id"])["response"]
    assert (
        next(c for c in w["checks"] if c["id"] == "review:" + comment["id"])["status"]
        == "fail"
    )
    assert not w["release"]["ready"]
    response = post(
        client,
        f"/{w['order']['id']}/commands",
        {
            "expected_version": w["order"]["version"],
            "prompt": "Why is the short-circuit rating provisional?",
        },
    )
    assert response.status_code == 200 and response.json()["source_ids"]
    response = post(
        client,
        f"/{w['order']['id']}/commands",
        {
            "expected_version": w["order"]["version"],
            "prompt": "Draft a transmittal to the reviewer.",
        },
    )
    assert response.status_code == 503 and "not configured" in response.text


def test_uploaded_pdf_replay_auto_checks_aliases_annotations_and_false_fix(client):
    from pathlib import Path

    from backend.worker import claim, process

    replay = Path(__file__).resolve().parents[1] / "examples" / "order-replay"
    w = create(client, "Uploaded PDF replay")

    def upload_file(filename):
        response = client.post(
            f"/api/projects/{w['order']['project_id']}/documents",
            files={
                "file": (filename, (replay / filename).read_bytes(), "application/pdf")
            },
            data={"kind": "auto"},
            headers={"Idempotency-Key": str(uuid4())},
        )
        assert response.status_code == 202, response.text
        process(claim())
        return client.get("/api/orders/" + w["order"]["id"]).json()

    for filename in [
        "01-project-specification.pdf",
        "02-customer-purchase-order.pdf",
        "03-accepted-exception.pdf",
        "04-approval-drawing-Rev-B.pdf",
        "05-customer-markups-Rev-B.pdf",
    ]:
        w = upload_file(filename)
    # Identity/authority decisions survive restaging when another one is accepted.
    for _ in range(4):
        action = next(
            (
                a
                for a in w["actions"]
                if a["status"] == "pending"
                and (a["type"] == "alias" or "accept_exception_document" in a["after"])
            ),
            None,
        )
        assert action, w["actions"]
        w = decision(client, w, action)
    assert len(w["comments"]) == 11
    assert len(w["ledger"]) == 34
    assert sum(r["status"] == "fail" for r in w["ledger"]) == 7
    assert (
        next(r for r in w["ledger"] if r["attribute"] == "bus_material")["expected"]
        == "aluminum"
    )
    cloud = next(c for c in w["checks"] if c["id"].startswith("source-review:"))
    assert cloud["status"] == "unknown"
    comment7 = next(c for c in w["comments"] if c["number"] == "7")
    span = next(sp for sp in w["sources"] if sp["id"] == comment7["source_ids"][0])
    assert span["location"]["annotation_type"] == "FreeText"
    assert len(span["location"]["bbox"]) == 4
    version_b = w["order"]["version"]
    w = upload_file("04-approval-drawing-Rev-C.pdf")
    assert len(w["ledger"]) == 34
    assert sum(r["status"] == "fail" for r in w["ledger"]) == 1
    assert (
        next(r for r in w["ledger"] if r["attribute"] == "short_circuit_ka")["actual"]
        == 65
    )
    assert (
        next(c for c in w["checks"] if c["id"] == "review:" + comment7["id"])["status"]
        == "fail"
    )
    comment8 = next(c for c in w["comments"] if c["number"] == "8")
    assert (
        next(c for c in w["checks"] if c["id"] == "review:" + comment8["id"])["status"]
        == "conflict"
    )
    result = client.get(
        f"/api/orders/{w['order']['id']}/diff?from_version={version_b}"
    ).json()
    assert sum(c["before"] != c["after"] for c in result["changes"]) == 6
    assert not w["release"]["ready"]
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get(
        f"/api/orders/{w['order']['id']}/exports/response-matrix?format=xlsx"
    )
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames
    response = client.get(
        f"/api/orders/{w['order']['id']}/exports/bom-delta?from_version={version_b}"
    )
    assert response.status_code == 200
    assert load_workbook(BytesIO(response.content)).sheetnames


def test_named_source_region_review_can_be_signed_but_reopens_on_changed_file(client):
    w = fixture(client, actual=85)
    document_id, _ = evidence(
        w, "Review markups.txt", "markup", ["Comment 1: MSB-01 | short_circuit_ka: 85"]
    )
    with Session.begin() as s:
        document = get(s, Document, document_id)
        document.coverage = [
            {
                "label": "Cloud 1",
                "state": "unsupported",
                "detail": "Empty cloud annotation requires visual review.",
            }
        ]
    w = refresh(client, w)
    check = next(c for c in w["checks"] if c["id"].startswith("source-review:"))
    response = post(
        client,
        f"/{w['order']['id']}/waivers",
        {
            "expected_version": w["order"]["version"],
            "check_id": check["id"],
            "fingerprint": check["fingerprint"],
            "signer": "Test Engineer",
            "reason": "Opened original page one; cloud repeats comment 1, now verified against the drawing.",
        },
    )
    assert response.status_code == 200, response.text
    w = response.json()
    assert next(c for c in w["checks"] if c["id"] == check["id"])["status"] == "waived"
    assert not w["release"]["ready"]  # response still must be accepted


def test_model_can_only_stage_evidence_backed_changes(client, monkeypatch):
    from backend.orders import commands

    w = fixture(client)
    rating = next(r for r in w["ledger"] if r["attribute"] == "short_circuit_ka")

    def proposal(prompt, data):
        return {
            "answer": "Prepared a change for review.",
            "source_ids": rating["source_ids"],
            "action": {
                "type": "configuration",
                "title": "Bring SCCR in line with the PO",
                "summary": "Review this supported change.",
                "after": {
                    "device": "MSB-01",
                    "attribute": "short_circuit_ka",
                    "value": 85,
                    "unit": "kA",
                },
                "source_ids": rating["source_ids"],
            },
        }

    monkeypatch.setattr(commands, "model_proposal", proposal)
    response = post(
        client,
        f"/{w['order']['id']}/commands",
        {
            "expected_version": w["order"]["version"],
            "prompt": "Prepare the necessary engineering change.",
        },
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["actions"][0]["status"] == "pending"
    assert (
        next(
            r
            for r in result["workspace"]["ledger"]
            if r["attribute"] == "short_circuit_ka"
        )["actual"]
        == 65
    )
    assert not result["workspace"]["release"]["approval"]
    saved_version = result["workspace"]["order"]["version"]

    def invented(prompt, data):
        result = proposal(prompt, data)
        result["action"]["after"]["value"] = 999
        return result

    monkeypatch.setattr(commands, "model_proposal", invented)
    response = post(
        client,
        f"/{w['order']['id']}/commands",
        {
            "expected_version": saved_version,
            "prompt": "Prepare a change from this invented value.",
        },
    )
    assert response.status_code == 422 and "not supported" in response.text
    assert (
        client.get("/api/orders/" + w["order"]["id"]).json()["order"]["version"]
        == saved_version
    )


def test_parallel_decisions_use_one_current_order_version(client):
    from concurrent.futures import ThreadPoolExecutor

    w = fixture(client)
    action = next(a for a in w["actions"] if a["status"] == "pending")
    url = f"/{w['order']['id']}/actions/{action['id']}/accept"
    body = {
        "expected_version": w["order"]["version"],
        "reason": "Reviewed exact source evidence.",
    }
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: post(client, url, body), range(2)))
    assert sorted(result.status_code for result in results) == [200, 409]
    current = client.get("/api/orders/" + w["order"]["id"]).json()
    assert current["order"]["version"] == w["order"]["version"] + 1
    assert len(current["tasks"]) == 4


def test_legacy_quote_views_exclude_order_backing_projects(client):
    from test_workflow import project

    legacy = project(client)
    order = create(client)
    projects = client.get("/api/projects").json()
    assert [p["id"] for p in projects] == [legacy["id"]]
    queue = client.get("/api/work-queue").json()
    assert queue["items"] and all(
        item["project_id"] == legacy["id"] for item in queue["items"]
    )
    assert all(
        item["project_id"] != order["order"]["project_id"] for item in queue["items"]
    )
    # The immutable source backing remains directly addressable for ingestion.
    assert (
        client.get("/api/projects/" + order["order"]["project_id"]).status_code == 200
    )
    assert [item["id"] for item in client.get("/api/orders").json()["orders"]] == [
        order["order"]["id"]
    ]


def test_history_processing_state_and_bom_export_selected_target_revision(client):
    from io import BytesIO

    from openpyxl import load_workbook

    w = fixture(client)
    before = w["order"]["version"]
    evidence(
        w,
        "Approval drawing Rev C.txt",
        "drawing",
        [("short_circuit_ka", 85), ("bus_material", "copper")],
    )
    w = refresh(client, w)
    selected = w["order"]["version"]
    evidence(
        w,
        "Approval drawing Rev D.txt",
        "drawing",
        [("short_circuit_ka", 100), ("bus_material", "copper")],
    )
    w = refresh(client, w)
    evidence(w, "Pending drawing.txt", "drawing", [], state="queued")
    w = refresh(client, w)
    assert w["history"][0]["is_processing"] is True
    assert (
        next(h for h in w["history"] if h["version"] == selected)["is_processing"]
        is False
    )
    response = client.get(
        f"/api/orders/{w['order']['id']}/exports/bom-delta?from_version={before}&to_version={selected}"
    )
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.cell(4, 2).value == "Rev C"
    assert sheet.cell(4, 4).value == selected
    rating = next(
        row
        for row in sheet.iter_rows(values_only=True)
        if row[1] == "Short-circuit rating"
    )
    assert str(rating[2]) == "65" and str(rating[3]) == "85"
    assert f"-v{selected}-" in response.headers["content-disposition"]


def test_document_classification_corrects_unknown_role_without_changing_original(
    client,
):
    from backend.domain.service import scoped
    from backend.storage.models import Span

    w = fixture(client, actual=85)
    document_id, _ids = evidence(
        w,
        "Untitled customer attachment.txt",
        "unknown",
        [("short_circuit_ka", 85), ("bus_material", "copper")],
    )
    w = refresh(client, w)
    assert (
        next(d for d in w["documents"] if d["id"] == document_id)["role"] == "unknown"
    )
    assert not w["release"]["ready"]
    with Session() as s:
        document = get(s, Document, document_id)
        original = (document.blob, document.sha256)
        source_values = [
            (sp.id, sp.text, sp.location)
            for sp in s.scalars(scoped(Span).where(Span.document_id == document_id))
        ]
    body = {
        "expected_version": w["order"]["version"],
        "role": "drawing",
        "revision_label": "C",
    }
    key = str(uuid4())
    response = post(
        client, f"/{w['order']['id']}/documents/{document_id}/classify", body, key
    )
    assert response.status_code == 200, response.text
    current = response.json()
    doc = next(d for d in current["documents"] if d["id"] == document_id)
    assert doc["role"] == "drawing" and doc["revision_label"] == "C"
    assert not doc["warnings"] and current["release"]["ready"]
    assert current["order"]["version"] == w["order"]["version"] + 1
    assert current["events"][0]["kind"] == "classify"
    assert current["events"][0]["data"]["before"]["order_role"] == "unknown"
    assert (
        post(
            client, f"/{w['order']['id']}/documents/{document_id}/classify", body, key
        ).json()
        == current
    )
    with Session() as s:
        document = get(s, Document, document_id)
        assert (document.blob, document.sha256) == original
        assert [
            (sp.id, sp.text, sp.location)
            for sp in s.scalars(scoped(Span).where(Span.document_id == document_id))
        ] == source_values
    release_body = {
        "expected_version": current["order"]["version"],
        "snapshot_hash": current["order"]["snapshot_hash"],
        "signer": "Test Engineer",
        "reason": "Evidence reviewed for exact release package.",
    }
    response = post(client, f"/{w['order']['id']}/release", release_body)
    assert response.status_code == 200, response.text
    response = post(
        client,
        f"/{w['order']['id']}/documents/{document_id}/classify",
        {
            "expected_version": current["order"]["version"],
            "role": "drawing",
            "revision_label": "D",
        },
    )
    assert (
        response.status_code == 200 and response.json()["release"]["approval"] is None
    )


def test_document_classification_scopes_and_waits_for_source_reads(client):
    w, other = fixture(client), fixture(client)
    foreign = other["documents"][0]["id"]
    response = post(
        client,
        f"/{w['order']['id']}/documents/{foreign}/classify",
        {"expected_version": w["order"]["version"], "role": "drawing"},
    )
    assert response.status_code == 403
    document_id, _ = evidence(w, "Reading drawing.txt", "drawing", [], state="queued")
    w = refresh(client, w)
    response = post(
        client,
        f"/{w['order']['id']}/documents/{document_id}/classify",
        {
            "expected_version": w["order"]["version"],
            "role": "drawing",
            "revision_label": "C",
        },
    )
    assert response.status_code == 409 and "finishes reading" in response.text
    response = post(
        client,
        f"/{w['order']['id']}/documents/{document_id}/classify",
        {"expected_version": w["order"]["version"], "role": "invented_authority"},
    )
    assert response.status_code == 422


def test_changed_since_named_revision_skips_intermediate_processing_snapshot(client):
    w = fixture(client)
    completed_b = w["order"]["version"]
    evidence(w, "Approval drawing Rev C.txt", "drawing", [], state="queued")
    w = refresh(client, w)
    assert (
        w["history"][0]["revision_label"] == "Rev B"
        and w["history"][0]["is_processing"]
    )
    evidence(
        w,
        "Approval drawing Rev C.txt",
        "drawing",
        [("short_circuit_ka", 85), ("bus_material", "copper")],
    )
    w = refresh(client, w)
    response = post(
        client,
        f"/{w['order']['id']}/commands",
        {
            "expected_version": w["order"]["version"],
            "prompt": "What changed since Rev B?",
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["answer"].startswith(f"Compared revision {completed_b} with")
    assert "65 → 85" in response.json()["answer"]


def test_unseeded_orders_reproduce_full_workflow_with_independent_inputs(client):
    """Exercise public intake/worker/decisions without the demo or DB-insert helpers."""
    from backend.worker import claim, process

    def upload(w, role, text, revision=""):
        response = client.post(
            f"/api/projects/{w['order']['project_id']}/documents",
            files={"file": (f"{role}.txt", text.encode(), "text/plain")},
            data={"kind": role, "revision_label": revision},
            headers={"Idempotency-Key": str(uuid4())},
        )
        assert response.status_code == 202, response.text
        process(claim())
        return client.get(f"/api/orders/{w['order']['id']}").json()

    completed = []
    for title, tag, required, initial in [
        ("Harbor expansion", "SWGR-27", 100, 50),
        ("Research facility", "LINEUP-X9", 42, 25),
    ]:
        response = post(client, "", {"title": title, "customer": title + " team"})
        assert response.status_code == 200
        w = response.json()
        assert not w["order"]["synthetic"]
        assert not w["ledger"]
        w = upload(w, "purchase_order", f"{tag} | SCCR: {required} kA")
        w = upload(w, "approval_drawing", f"{tag} | SCCR: {initial} kA", "03")
        row = next(r for r in w["ledger"] if r["attribute"] == "short_circuit_ka")
        assert (row["device"], row["expected"], row["actual"], row["status"]) == (
            tag, required, initial, "fail"
        )
        baseline = w["order"]["version"]
        action = next(a for a in w["actions"] if a["status"] == "pending"
                      and a["after"].get("attribute") == "short_circuit_ka")
        w = decision(client, w, action)
        assert not w["release"]["ready"]
        assert len(w["tasks"]) == 4
        w = upload(w, "approval_drawing", f"{tag} | SCCR: {required} kA", "04")
        for role in ("bom", "supplier_po", "nameplate"):
            w = upload(w, role, f"{tag} | SCCR: {required} kA", "04")
        assert all(t["status"] == "verified" for t in w["tasks"])
        assert w["release"]["ready"], w["release"]
        diff = client.get(f"/api/orders/{w['order']['id']}/diff?from_version={baseline}")
        assert diff.status_code == 200
        assert tag in diff.text
        released = post(client, f"/{w['order']['id']}/release", {
            "expected_version": w["order"]["version"],
            "signer": "Test reviewer", "reason": "Checked uploaded source documents.",
            "snapshot_hash": w["order"]["snapshot_hash"],
        })
        assert released.status_code == 200, released.text
        assert released.json()["release"]["approval"]
        completed.append(released.json())
    # A second order's uploads and decisions must not change the first order.
    first = client.get(f"/api/orders/{completed[0]['order']['id']}").json()
    assert first["order"]["version"] == completed[0]["order"]["version"]
    assert first["ledger"] == completed[0]["ledger"]
    assert first["release"]["approval"] == completed[0]["release"]["approval"]
