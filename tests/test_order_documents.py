from io import BytesIO
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

from backend.orders.exports import (
    bom_delta_xlsx,
    response_matrix_html,
    response_matrix_xlsx,
)
from backend.orders.extraction import (
    ATTRIBUTE_REGISTRY,
    extract_order_document,
    normalize_value,
    parse_pdf_annotations,
)

REPLAY = Path(__file__).resolve().parents[1] / "examples" / "order-replay"


def read_fixture(filename):
    path = REPLAY / filename
    spans = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            for line_number, line in enumerate(
                (page.extract_text() or "").splitlines(), 1
            ):
                spans.append(
                    {
                        "id": f"{filename}:{page_number}:{line_number}",
                        "text": line,
                        "location": {"page": page_number},
                    }
                )
    for annotation in parse_pdf_annotations(path):
        if annotation["text"]:
            spans.append(
                {
                    "id": annotation["id"],
                    "text": annotation["text"],
                    "location": {
                        "page": annotation["page"],
                        "bbox": annotation["bbox"],
                        "annotation_id": annotation["id"],
                        "annotation_type": annotation["kind"],
                    },
                }
            )
    return extract_order_document({"name": filename, "meta": {}, "coverage": []}, spans)


def test_order_reads_vector_pdf_facts_and_explicit_aliases():
    result = read_fixture("01-project-specification.pdf")
    assert result["role"] == "specification"
    assert len(result["facts"]) == len(ATTRIBUTE_REGISTRY) == 34
    assert {item["alias"]: item["device"] for item in result["aliases"]} == {
        "MSB Main": "MSB-01",
        "52-M1": "MSB-01",
        "BOM line 3": "MSB-01",
    }
    assert all(item["source_ids"] for item in result["facts"])
    assert (
        next(
            item for item in result["facts"] if item["attribute"] == "rated_voltage_v"
        )["value"]
        == 480
    )


def test_replay_revision_c_preserves_claimed_fixed_failure():
    before = read_fixture("04-approval-drawing-Rev-B.pdf")
    after = read_fixture("04-approval-drawing-Rev-C.pdf")
    assert before["revision_label"] == "B"
    assert after["revision_label"] == "C"
    values_b = {item["attribute"]: item["value"] for item in before["facts"]}
    values_c = {item["attribute"]: item["value"] for item in after["facts"]}
    assert len(values_b) == len(values_c) == 34
    assert values_b["short_circuit_ka"] == values_c["short_circuit_ka"] == 65
    assert values_b["main_breaker_trip_a"] == 2500
    assert values_c["main_breaker_trip_a"] == 3000
    assert {key for key in values_b if values_b[key] != values_c[key]} == {
        "main_breaker_trip_a",
        "enclosure_type",
        "communications_protocol",
        "cable_entry",
        "enclosure_width_mm",
        "nameplate_text",
    }


def test_order_annotation_reads_numbered_callout_and_does_not_invent_cloud_content():
    annotations = parse_pdf_annotations(REPLAY / "05-customer-markups-Rev-B.pdf")
    callout, cloud = annotations
    assert callout["kind"] == "FreeText"
    assert callout["text"].startswith("Comment 7")
    assert all(0 <= value <= 1 for value in callout["bbox"])
    assert callout["bbox"][0] < callout["bbox"][2]
    assert callout["bbox"][1] < callout["bbox"][3]
    assert cloud["kind"] == "Polygon"
    assert cloud["text"] == "" and cloud["needs_review"]
    extracted = read_fixture("05-customer-markups-Rev-B.pdf")
    assert [item["number"] for item in extracted["comments"]] == [
        str(number) for number in range(1, 12)
    ]
    comment = extracted["comments"][6]
    assert comment["required_value"] == 85
    assert comment["source_ids"] == ["review-comment-7"]
    assert extracted["facts"] == []
    assert extracted["comments"][-1]["attribute"] is None


def test_order_number_units_are_normalized_without_guessing_ranges():
    import pytest

    assert normalize_value("rated_voltage_v", "0.48 kV") == (480, "V")
    assert normalize_value("short_circuit_ka", "85,000 A") == (85, "kA")
    assert normalize_value("enclosure_width_mm", '10"') == (254, "mm")
    assert normalize_value("surge_protection", "not included") == (False, None)
    for value in ["480/277 V", "65-85 kA", "85 kA provisional", "TBD"]:
        with pytest.raises(ValueError):
            normalize_value("short_circuit_ka", value)
    with pytest.raises(ValueError):
        normalize_value("quantity", "1.5")


def test_order_unknown_and_scanned_evidence_stays_unknown():
    extracted = extract_order_document(
        {
            "name": "unclassified.pdf",
            "coverage": [
                {"label": "Page 1", "state": "scanned", "detail": "No extractable text"}
            ],
        },
        [],
    )
    assert extracted["role"] == "unknown"
    assert extracted["facts"] == []
    assert any("No extractable text" in warning for warning in extracted["warnings"])
    no_device = extract_order_document(
        {"name": "spec.pdf"},
        [{"id": "s1", "text": "Rated voltage: 480 V", "location": {"page": 1}}],
    )
    assert no_device["facts"] == []
    assert "no explicit equipment tag" in no_device["warnings"][0]


def test_order_multiline_comments_keep_all_evidence():
    extracted = extract_order_document(
        {"name": "markups.pdf"},
        [
            {
                "id": "s1",
                "text": "1. MSB-01 | main_bus_a: 3200 A",
                "location": {"page": 1},
            },
            {
                "id": "s2",
                "text": "Coordinate with the incoming feeder.",
                "location": {"page": 1},
            },
            {
                "id": "s3",
                "text": "2. MSB-01 | Review the unexplained detail.",
                "location": {"page": 1},
            },
        ],
    )
    assert len(extracted["comments"]) == 2
    assert extracted["comments"][0]["source_ids"] == ["s1", "s2"]
    assert "Coordinate" in extracted["comments"][0]["text"]
    assert extracted["facts"] == []


def test_order_exports_use_accepted_responses_and_allowlisted_configuration_only():
    data = {
        "order": {
            "title": "Fictional <order>",
            "number": "DEMO-4512",
            "synthetic": True,
        },
        "comments": [
            {
                "number": "1",
                "device": "MSB-01",
                "text": "Confirm bus",
                "status": "open",
                "response": "",
                "draft_response": "SECRET DRAFT",
            },
            {
                "number": "2",
                "text": "Confirm enclosure",
                "status": "addressed",
                "response": "=not-a-formula",
            },
        ],
        "diff": {
            "changes": [
                {
                    "device": "MSB-01",
                    "attribute": "main_breaker_trip_a",
                    "before": 2500,
                    "after": 3000,
                    "source_ids": [],
                },
                {
                    "device": "MSB-01",
                    "attribute": "cost",
                    "before": 90000,
                    "after": 110000,
                },
            ]
        },
    }
    html = response_matrix_html(data)
    assert "SECRET DRAFT" not in html
    assert "Fictional &lt;order&gt;" in html
    assert "Awaiting an accepted response" in html
    matrix = load_workbook(BytesIO(response_matrix_xlsx(data)), data_only=False)
    cells = [cell.value for row in matrix.active for cell in row]
    assert "SECRET DRAFT" not in cells
    assert "'=not-a-formula" in cells
    delta = load_workbook(BytesIO(bom_delta_xlsx(data)), data_only=False)
    cells = [cell.value for row in delta.active for cell in row]
    assert 110000 not in cells and "cost" not in cells
    assert "Main breaker trip" in cells and "3000" in cells


def test_actual_pdf_parser_keeps_callout_sources_and_unreadable_cloud_warning():
    from types import SimpleNamespace
    from unittest.mock import patch

    from backend.ingestion import parser

    class CapturingSession:
        def __init__(self):
            self.spans = []

        def add(self, span):
            self.spans.append(span)

        def flush(self):
            pass

    for filename in (
        "01-project-specification.pdf",
        "04-approval-drawing-Rev-B.pdf",
        "05-customer-markups-Rev-B.pdf",
        "04-approval-drawing-Rev-C.pdf",
    ):
        document = SimpleNamespace(
            id="fixture-document",
            name=filename,
            blob=filename,
            kind="auto",
            meta={},
            coverage=[],
            state="queued",
        )
        session = CapturingSession()
        with patch.object(parser, "BLOBS", REPLAY):
            parser.parse(session, document)
        extracted = extract_order_document(document, session.spans)
        if "markups" in filename:
            assert len(extracted["comments"]) == 11
            comment = next(
                item for item in extracted["comments"] if item["number"] == "7"
            )
            source = next(
                span for span in session.spans if span.id == comment["source_ids"][0]
            )
            assert source.location["annotation_id"] == "review-comment-7"
            assert source.location["approximate"] is False
            assert any("graphical annotation" in item for item in extracted["warnings"])
            assert (
                len(
                    [
                        part
                        for part in document.coverage
                        if part["state"] == "unsupported"
                    ]
                )
                == 1
            )
        else:
            assert len(extracted["facts"]) == 34
            assert all(fact["location"].get("bbox") for fact in extracted["facts"])


def test_response_matrix_separates_approved_words_from_actual_check_results():
    from backend.orders.exports import response_matrix_rows

    rows = response_matrix_rows(
        {
            "comments": [
                {"id": "c1", "number": "1", "status": "open", "response": ""},
                {
                    "id": "c7",
                    "number": "7",
                    "status": "responded",
                    "response": "Fixed in Rev C.",
                },
                {
                    "id": "c8",
                    "number": "8",
                    "status": "responded",
                    "response": "Holding the exception.",
                },
            ],
            "checks": [
                {"id": "review:c1", "status": "pass"},
                {"id": "review:c7", "status": "fail"},
                {"id": "review:c8", "status": "conflict"},
            ],
        }
    )
    assert [row["status"] for row in rows] == [
        "Verified",
        "Not addressed",
        "Decision required",
    ]
    assert rows[0]["response"] == ""
    assert rows[1]["response"] == "Fixed in Rev C."


def test_multiple_comments_inside_one_annotation_are_split_without_losing_evidence():
    result = extract_order_document(
        {"name": "markups.pdf"},
        [
            {
                "id": "annotation-span",
                "text": "Comment 7 | MSB-01 | short_circuit_ka: 85 kA\nMatch the customer purchase order.\nComment 8 | MSB-01 | bus_material: copper",
                "location": {
                    "page": 1,
                    "annotation_id": "callout-1",
                    "annotation_type": "FreeText",
                },
            }
        ],
    )
    assert [item["number"] for item in result["comments"]] == ["7", "8"]
    assert result["comments"][0]["required_value"] == 85
    assert "Match the customer" in result["comments"][0]["text"]
    assert result["comments"][0]["source_ids"] == ["annotation-span"]
    assert result["comments"][1]["source_ids"] == ["annotation-span"]
